from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import os
from pathlib import Path

from fastapi import HTTPException, status


@dataclass(frozen=True)
class ExtractedTextBlock:
    text: str
    page: int | None = None


SUPPORTED_UPLOAD_SUFFIXES = {
    ".txt": "text_file",
    ".md": "markdown_file",
    ".markdown": "markdown_file",
    ".pdf": "pdf_file",
    ".pptx": "pptx_file",
}


def source_type_for_filename(filename: str) -> str:
    suffix = Path(filename.strip().lower()).suffix
    source_type = SUPPORTED_UPLOAD_SUFFIXES.get(suffix)
    if source_type:
        return source_type
    raise HTTPException(
        status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
        detail="Only .txt, .md, .markdown, .pdf, and .pptx uploads are supported",
    )


def extract_upload_text(filename: str, content: bytes) -> list[ExtractedTextBlock]:
    source_type = source_type_for_filename(filename)
    if source_type in {"text_file", "markdown_file"}:
        return [ExtractedTextBlock(text=decode_utf8_text(content))]
    if source_type == "pdf_file":
        return extract_pdf_text(content)
    if source_type == "pptx_file":
        return extract_pptx_text(content)
    raise AssertionError(f"Unhandled source type: {source_type}")


def decode_utf8_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded TXT/MD file must be UTF-8 encoded",
        ) from exc


def extract_pdf_text(content: bytes) -> list[ExtractedTextBlock]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="PDF extraction dependency is not installed",
        ) from exc

    try:
        reader = PdfReader(BytesIO(content))
        blocks = [
            ExtractedTextBlock(text=text, page=index)
            for index, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "").strip())
        ]
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded PDF could not be read",
        ) from exc

    if blocks:
        return blocks
    return extract_pdf_ocr_text(content)


def extract_pdf_ocr_text(content: bytes) -> list[ExtractedTextBlock]:
    if os.environ.get("OCR_ENABLED", "").strip().lower() not in {"1", "true", "yes"}:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=(
                "No extractable text found in PDF. Enable OCR with OCR_ENABLED=1 "
                "and install OCR dependencies for scanned PDFs."
            ),
        )

    try:
        from pdf2image import convert_from_bytes
        import pytesseract
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="OCR dependencies are not installed",
        ) from exc

    try:
        images = convert_from_bytes(content)
        blocks = [
            ExtractedTextBlock(text=text, page=index)
            for index, image in enumerate(images, start=1)
            if (text := pytesseract.image_to_string(image).strip())
        ]
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="OCR failed for uploaded PDF",
        ) from exc

    if not blocks:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="OCR completed but found no text in uploaded PDF",
        )
    return blocks


def extract_pptx_text(content: bytes) -> list[ExtractedTextBlock]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="PPTX extraction dependency is not installed",
        ) from exc

    try:
        presentation = Presentation(BytesIO(content))
        blocks = []
        for slide in presentation.slides:
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if texts:
                blocks.append(ExtractedTextBlock(text="\n".join(texts)))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded PPTX could not be read",
        ) from exc

    if not blocks:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="No extractable text found in PPTX",
        )
    return blocks
