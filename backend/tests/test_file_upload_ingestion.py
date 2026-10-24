import asyncio
import importlib
from io import BytesIO
import sys
from pathlib import Path

from fastapi import HTTPException
from httpx import ASGITransport, AsyncClient
import pytest


def load_app(tmp_path, monkeypatch, name: str):
    monkeypatch.setenv("NOTEBOOKLM_DB_PATH", str(tmp_path / f"{name}.db"))
    backend_dir = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(backend_dir))

    main = importlib.import_module("app.main")
    database = importlib.import_module("app.database")
    extractors = importlib.import_module("app.extractors")
    limits = importlib.import_module("app.limits")
    services = importlib.import_module("app.services")
    database.init_db()
    return main, extractors, limits, services


def test_upload_pdf_preserves_page_metadata(tmp_path, monkeypatch):
    asyncio.run(run_pdf_upload_test(tmp_path, monkeypatch))


async def run_pdf_upload_test(tmp_path, monkeypatch):
    main, extractors, _limits, services = load_app(tmp_path, monkeypatch, "pdf")

    def fake_extract(filename: str, content: bytes):
        assert filename == "lecture.pdf"
        assert content == b"%PDF fake"
        return [
            extractors.ExtractedTextBlock(
                text="Lecture PDF explains integrals and derivatives.",
                page=2,
            )
        ]

    monkeypatch.setattr(services, "extract_upload_text", fake_extract)

    transport = ASGITransport(app=main.app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        notebook_response = await client.post("/api/notebooks", json={"title": "Math"})
        notebook = notebook_response.json()

        upload_response = await client.post(
            f"/api/notebooks/{notebook['id']}/sources/upload",
            params={"filename": "lecture.pdf"},
            content=b"%PDF fake",
            headers={"content-type": "application/pdf"},
        )
        assert upload_response.status_code == 201
        source = upload_response.json()
        assert source["type"] == "pdf_file"

        chat_response = await client.post(
            f"/api/notebooks/{notebook['id']}/chat",
            json={"question": "What explains integrals?"},
        )
        assert chat_response.status_code == 200
        citation = chat_response.json()["citations"][0]
        assert citation["page"] == 2

        evidence_response = await client.get(f"/api/chunks/{citation['chunk_id']}")
        assert evidence_response.status_code == 200
        evidence = evidence_response.json()
        assert evidence["page"] == 2
        assert "integrals" in evidence["text"]


def test_upload_pptx_has_no_page_metadata(tmp_path, monkeypatch):
    asyncio.run(run_pptx_upload_test(tmp_path, monkeypatch))


async def run_pptx_upload_test(tmp_path, monkeypatch):
    main, extractors, _limits, services = load_app(tmp_path, monkeypatch, "pptx")

    def fake_extract(filename: str, content: bytes):
        assert filename == "slides.pptx"
        return [
            extractors.ExtractedTextBlock(
                text="Slide deck covers graphs, breadth first search, and queues.",
            )
        ]

    monkeypatch.setattr(services, "extract_upload_text", fake_extract)

    transport = ASGITransport(app=main.app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        notebook_response = await client.post(
            "/api/notebooks",
            json={"title": "Algorithms"},
        )
        notebook = notebook_response.json()

        upload_response = await client.post(
            f"/api/notebooks/{notebook['id']}/sources/upload",
            params={"filename": "slides.pptx"},
            content=b"pptx fake",
            headers={
                "content-type": (
                    "application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation"
                )
            },
        )
        assert upload_response.status_code == 201
        source = upload_response.json()
        assert source["type"] == "pptx_file"

        chat_response = await client.post(
            f"/api/notebooks/{notebook['id']}/chat",
            json={"question": "What covers breadth first search?"},
        )
        citation = chat_response.json()["citations"][0]
        assert citation["page"] is None


def test_legacy_ppt_and_oversized_uploads_are_rejected(tmp_path, monkeypatch):
    asyncio.run(run_rejected_upload_test(tmp_path, monkeypatch))


async def run_rejected_upload_test(tmp_path, monkeypatch):
    main, _extractors, limits, _services = load_app(tmp_path, monkeypatch, "reject")

    transport = ASGITransport(app=main.app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        notebook_response = await client.post(
            "/api/notebooks",
            json={"title": "Rejected uploads"},
        )
        notebook = notebook_response.json()

        ppt_response = await client.post(
            f"/api/notebooks/{notebook['id']}/sources/upload",
            params={"filename": "legacy.ppt"},
            content=b"old powerpoint",
            headers={"content-type": "application/octet-stream"},
        )
        assert ppt_response.status_code == 415

        oversized_response = await client.post(
            f"/api/notebooks/{notebook['id']}/sources/upload",
            params={"filename": "too-large.txt"},
            content=b"x" * (limits.MAX_UPLOAD_BYTES + 1),
            headers={"content-type": "application/octet-stream"},
        )
        assert oversized_response.status_code == 413


def test_pdf_ocr_disabled_reports_clear_error(tmp_path, monkeypatch):
    _main, extractors, _limits, _services = load_app(tmp_path, monkeypatch, "ocr")
    monkeypatch.delenv("OCR_ENABLED", raising=False)

    with pytest.raises(HTTPException) as exc_info:
        extractors.extract_pdf_ocr_text(b"image only")

    assert exc_info.value.status_code == 422
    assert "Enable OCR" in exc_info.value.detail


def test_real_pdf_extractor_keeps_page_number(tmp_path, monkeypatch):
    _main, extractors, _limits, _services = load_app(tmp_path, monkeypatch, "real-pdf")

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 18 Tf 72 720 Td (Smoke PDF mentions eigenvalues.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)

    blocks = extractors.extract_pdf_text(buffer.getvalue())

    assert blocks[0].page == 1
    assert "eigenvalues" in blocks[0].text


def test_real_pptx_extractor_reads_slide_text(tmp_path, monkeypatch):
    _main, extractors, _limits, _services = load_app(tmp_path, monkeypatch, "real-pptx")

    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    text_box = slide.shapes.add_textbox(0, 0, 5_000_000, 500_000)
    text_box.text = "Smoke PPTX mentions Dijkstra shortest paths."
    buffer = BytesIO()
    presentation.save(buffer)

    blocks = extractors.extract_pptx_text(buffer.getvalue())

    assert blocks[0].page is None
    assert "Dijkstra" in blocks[0].text
